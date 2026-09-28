"""Deterministic resume text extraction — the ONLY place in this project
that reads a PDF/DOCX file's raw bytes. Everything downstream (structured
parsing, claims) only ever sees the plain text this module produces.

Two responsibilities, kept separate:
  - `classify_resume_file`: decide whether an upload is really a PDF or a
    DOCX at all (never trusts a client-supplied filename/content-type
    alone — see its docstring).
  - `extract_resume_text`: turn validated bytes of a known kind into plain
    text, or raise `ResumeExtractionError` with a message safe to show the
    candidate (never a raw library exception).

No OCR here (Phase 2 scope, see the task brief): a scanned PDF with no
text layer is a clear, reported extraction failure, never a silently
empty/invented profile.
"""

import io
import re
from typing import Literal

from docx import Document
from pypdf import PdfReader

from app.resume.exceptions import ResumeExtractionError, UnsupportedResumeFileTypeError

ResumeFileKind = Literal["pdf", "docx"]

_PDF_MAGIC = b"%PDF-"
# The ZIP local-file-header signature — DOCX (and every other OOXML
# format) is a ZIP archive, so this is necessary but not sufficient; see
# `_looks_like_docx_package` below for the OOXML-specific check on top of
# it.
_ZIP_MAGIC = b"PK\x03\x04"
_DOCX_PACKAGE_MARKER = b"word/document.xml"

# A resume that "extracted successfully" but yielded almost no real text
# (e.g. a scanned PDF with no text layer, or a handful of stray glyphs) is
# not usefully different from one that failed outright — treated the same
# way rather than persisted as a barely-populated profile.
_MIN_EXTRACTED_TEXT_CHARS = 20


def _looks_like_docx_package(data: bytes) -> bool:
    # A cheap, dependency-free stand-in for fully parsing the zip's
    # central directory: real DOCX files store `word/document.xml`
    # uncompressed-name-visible near the start of most producers' output,
    # so a raw substring scan over the first portion of the archive is
    # enough to distinguish "some other zip/OOXML file" from a Word
    # document without importing `zipfile` twice or opening it here (the
    # real open-and-parse happens once, in `_extract_docx_text`).
    return _DOCX_PACKAGE_MARKER in data[:65536]


def classify_resume_file(filename: str, declared_content_type: str, data: bytes) -> ResumeFileKind:
    """Decide whether `data` is really a PDF or DOCX resume.

    Never trusts the filename extension or the client-supplied
    `declared_content_type` alone — either can be wrong or deliberately
    spoofed (Content-Type is just a header the client sets; a filename is
    just a string). The actual file signature is authoritative: a PDF
    must start with the PDF magic bytes, and a DOCX must be a ZIP archive
    that actually contains `word/document.xml` (not just any ZIP-based
    format sharing the same outer magic bytes, e.g. a plain .zip or an
    .xlsx renamed to .docx). `declared_content_type` and the extension are
    used only to decide WHICH signature to check for — never to bypass
    the check.
    """
    name = (filename or "").strip().lower()
    declared = (declared_content_type or "").strip().lower()

    looks_like_pdf = name.endswith(".pdf") or "pdf" in declared
    looks_like_docx = name.endswith(".docx") or "wordprocessingml" in declared

    if looks_like_pdf and data.startswith(_PDF_MAGIC):
        return "pdf"
    if looks_like_docx and data.startswith(_ZIP_MAGIC) and _looks_like_docx_package(data):
        return "docx"

    # Also accept a file whose actual bytes are unambiguous even if the
    # filename/content-type hint was missing or wrong (e.g. a bare
    # "resume" with no extension) — the signature is what actually
    # matters, not the hint.
    if data.startswith(_PDF_MAGIC):
        return "pdf"
    if data.startswith(_ZIP_MAGIC) and _looks_like_docx_package(data):
        return "docx"

    raise UnsupportedResumeFileTypeError(
        "Only PDF and DOCX resumes are supported, and this file's contents don't match either format."
    )


def _normalize_whitespace(text: str) -> str:
    # Collapse runs of blank lines and trailing/leading whitespace per
    # line — a common PDF/DOCX extraction artifact — without touching the
    # actual wording, so the original content is preserved verbatim.
    lines = [line.strip() for line in text.splitlines()]
    collapsed: list[str] = []
    blank_run = 0
    for line in lines:
        if line:
            collapsed.append(line)
            blank_run = 0
        else:
            blank_run += 1
            if blank_run <= 1:
                collapsed.append("")
    normalized = "\n".join(collapsed).strip()
    # Some PDF extractors leave runs of internal spaces/tabs; a resume's
    # meaning never depends on exact inter-word spacing, so this is safe.
    return re.sub(r"[ \t]{2,}", " ", normalized)


def _extract_pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # An empty-password attempt is the only decryption this
            # project performs; anything actually password-protected is
            # reported the same as any other unreadable file rather than
            # guessed at.
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ResumeExtractionError("This PDF is password-protected and can't be read.") from exc
        pages_text = [page.extract_text() or "" for page in reader.pages]
    except ResumeExtractionError:
        raise
    except Exception as exc:
        raise ResumeExtractionError("This PDF could not be read — it may be corrupted.") from exc

    return "\n\n".join(pages_text)


def _extract_docx_text(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
        parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
        # Table text is included too (Step 3: "extract paragraphs and
        # relevant table text if practical") — section ordering is
        # otherwise preserved as python-docx yields it (document order).
        for table in document.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text for cell in row.cells)
                if row_text.strip():
                    parts.append(row_text)
    except Exception as exc:
        raise ResumeExtractionError("This DOCX could not be read — it may be corrupted.") from exc

    return "\n".join(parts)


def extract_resume_text(kind: ResumeFileKind, data: bytes) -> str:
    """Return normalized, non-empty plain text for a resume already
    classified by `classify_resume_file`. Raises `ResumeExtractionError`
    (with a message safe to show the candidate) for a corrupted file or
    one with no meaningful extractable text — e.g. a scanned PDF with no
    text layer — never a silently empty result passed on as if it
    succeeded.
    """
    raw_text = _extract_pdf_text(data) if kind == "pdf" else _extract_docx_text(data)
    normalized = _normalize_whitespace(raw_text)

    if len(re.sub(r"\s", "", normalized)) < _MIN_EXTRACTED_TEXT_CHARS:
        raise ResumeExtractionError(
            "No readable text was found in this file. If it's a scanned PDF, try uploading a "
            "text-based PDF or a DOCX instead."
        )

    return normalized
